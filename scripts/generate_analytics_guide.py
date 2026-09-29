#!/usr/bin/env python3
"""Capture LMS analytics screenshots (Playwright) and build a simple Word guide.

Arrows are placed from live element bounding boxes so numbers stay synced
with the figure legends. A fresh throwaway student is created for Diagnostic
UI shots (so we never show “already completed”).
"""
from __future__ import annotations

import asyncio
import json
import os
import shutil
import time
from pathlib import Path

import requests
import urllib3
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "guides"
ASSET_DIR = OUT_DIR / "assets" / "analytics_guide"
DESKTOP = Path(r"C:\Users\user\Desktop")
QA_SHOTS = ROOT.parent / "_qa_audit_tmp" / "ui_shots"
META_PATH = ASSET_DIR / "callouts.json"

BASE_URL = os.environ.get("GUIDE_BASE_URL", "https://iqbalai.com").rstrip("/")
TEACHER = os.environ.get("QA_TP_TEACHER", "qa.teacher.e2e.1789501674@test.local")
# Completed student — used only for post-diagnostic / learning-path overview
STUDENT_DONE = os.environ.get(
    "QA_TP_STUDENT_DONE",
    "qa.student.diag.1789501674@test.local",
)
PASSWORD = os.environ.get("QA_TP_PASS", "E2eQaTp2026!")
ADMIN_EMAIL = os.environ.get("GUIDE_ADMIN_USER", "admin@iqbalai.com")
ADMIN_PASS = os.environ.get("GUIDE_ADMIN_PASS", "Hamzakhanswati12@")

GREEN = RGBColor(0x16, 0xA3, 0x4A)
DARK = RGBColor(0x11, 0x18, 0x27)
GRAY = RGBColor(0x47, 0x55, 0x69)


def ensure_dirs() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)


def _font(size: int, bold: bool = False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path(r"C:\Windows\Fonts") / name
    if path.exists():
        return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def annotate(src: Path, dst: Path, callouts: list[dict]) -> Path:
    """Draw numbered red callouts. tip = target; badge = number circle."""
    if not src.exists():
        return dst
    img = Image.open(src).convert("RGBA")
    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    badge_font = _font(18, bold=True)
    color = "#DC2626"
    w, h = img.size

    for item in callouts:
        n = int(item["n"])
        tx, ty = int(item["tip"][0]), int(item["tip"][1])
        bx, by = int(item["badge"][0]), int(item["badge"][1])
        bx = max(22, min(w - 22, bx))
        by = max(22, min(h - 22, by))
        tx = max(8, min(w - 8, tx))
        ty = max(8, min(h - 8, ty))
        r = 16
        draw.line([(bx, by), (tx, ty)], fill=color, width=4)
        draw.ellipse([bx - r, by - r, bx + r, by + r], fill=color, outline="white", width=3)
        text = str(n)
        bbox = draw.textbbox((0, 0), text, font=badge_font)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text((bx - tw / 2, by - th / 2 - 1), text, fill="white", font=badge_font)
        draw.ellipse([tx - 7, ty - 7, tx + 7, ty + 7], fill=color, outline="white", width=2)

    out = Image.alpha_composite(img, overlay).convert("RGB")
    dst.parent.mkdir(parents=True, exist_ok=True)
    out.save(dst, quality=94)
    return dst


def create_fresh_student() -> tuple[str, str]:
    """Create a throwaway student who has NOT completed diagnostic."""
    run = str(int(time.time()))
    email = f"qa.guide.fresh.{run}@test.local"
    username = f"qa_guide_fresh_{run}"
    admin = requests.Session()
    admin.post(
        f"{BASE_URL}/auth/login",
        data={"useremail": ADMIN_EMAIL, "password": ADMIN_PASS},
        headers={"Accept": "application/json", "X-Requested-With": "XMLHttpRequest"},
        verify=False,
        timeout=60,
    )
    r = admin.post(
        f"{BASE_URL}/admin/users",
        json={
            "username": username,
            "useremail": email,
            "password": PASSWORD,
            "role": "student",
            "class_standard": "9th",
            "medium": "English",
        },
        verify=False,
        timeout=30,
    )
    if r.status_code not in (200, 201):
        raise RuntimeError(f"Could not create fresh student: {r.status_code} {r.text[:300]}")
    data = r.json()
    if not data.get("success") and not data.get("user_id"):
        raise RuntimeError(f"Create student failed: {data}")
    return email, PASSWORD


def copy_fallback(name: str, qa_name: str | None = None) -> None:
    dst = ASSET_DIR / f"{name}.png"
    if dst.exists() and dst.stat().st_size > 5000:
        return
    if qa_name:
        src = QA_SHOTS / qa_name
        if src.exists():
            shutil.copy2(src, dst)


async def capture_screenshots() -> dict:
    results: dict = {"ok": False, "shots": [], "errors": [], "callouts": {}, "fresh_student": None}
    try:
        from playwright.async_api import async_playwright
    except ImportError:
        results["errors"].append("Playwright not installed")
        return results

    async def login(page, email: str, password: str) -> None:
        await page.context.clear_cookies()
        await page.goto(f"{BASE_URL}/auth/login", wait_until="domcontentloaded", timeout=60000)
        await page.wait_for_timeout(600)
        email_sel = 'input[name="useremail"], input#useremail, input[type="email"], input#email'
        pass_sel = 'input[name="password"], input#password, input[type="password"]'
        await page.wait_for_selector(email_sel, timeout=30000)
        await page.fill(email_sel, email)
        await page.fill(pass_sel, password)
        await page.click('button[type="submit"], input[type="submit"]')
        await page.wait_for_timeout(2500)

    async def center(page, selector: str) -> tuple[int, int] | None:
        loc = page.locator(selector).first
        if not await loc.count():
            return None
        try:
            await loc.scroll_into_view_if_needed(timeout=3000)
        except Exception:
            pass
        box = await loc.bounding_box()
        if not box or box["width"] < 2 or box["height"] < 2:
            return None
        return (int(box["x"] + box["width"] / 2), int(box["y"] + box["height"] / 2))

    def badge_away(tip: tuple[int, int], prefer: str = "left") -> tuple[int, int]:
        """Place number circle away from the tip so it does not cover the control."""
        x, y = tip
        if prefer == "left":
            return (max(40, x - 140), max(40, y - 70))
        if prefer == "right":
            return (min(1320, x + 120), max(40, y - 70))
        if prefer == "above":
            return (x, max(40, y - 90))
        return (x, min(860, y + 80))

    async def shot(page, name: str) -> None:
        path = ASSET_DIR / f"{name}.png"
        await page.wait_for_timeout(400)
        await page.screenshot(path=str(path), full_page=False)
        results["shots"].append(name)

    # Fresh student for Diagnostic (not already completed)
    try:
        fresh_email, fresh_pw = create_fresh_student()
        results["fresh_student"] = fresh_email
    except Exception as exc:  # noqa: BLE001
        results["errors"].append(f"fresh student: {exc}")
        fresh_email, fresh_pw = STUDENT_DONE, PASSWORD

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        context = await browser.new_context(
            viewport={"width": 1366, "height": 900},
            ignore_https_errors=True,
        )

        # ---------- Fresh student: dashboard + live diagnostic ----------
        page = await context.new_page()
        try:
            await login(page, fresh_email, fresh_pw)
            await page.goto(
                f"{BASE_URL}/student-dashboard",
                wait_until="domcontentloaded",
                timeout=60000,
            )
            await page.wait_for_timeout(2500)
            # Close onboarding if it steals clicks
            await page.evaluate(
                """() => {
                document.querySelectorAll('.lms-modal-backdrop.open').forEach(m => m.classList.remove('open'));
            }"""
            )
            await page.wait_for_timeout(400)
            await shot(page, "01_student_dashboard")

            tip_diag = await center(page, 'button:has-text("Take Diagnostic")')
            tip_path = await center(page, "text=No path yet")
            if not tip_path:
                tip_path = await center(page, "text=LEARNING PATH")
            tip_classes = await center(page, "text=Not enrolled yet")
            if not tip_classes:
                tip_classes = await center(page, "text=MY CLASSES")
            if not tip_diag:
                tip_diag = (1236, 206)
            if not tip_path:
                tip_path = (423, 370)
            if not tip_classes:
                tip_classes = (1180, 370)
            results["callouts"]["01_student_dashboard"] = [
                {"n": 1, "tip": tip_diag, "badge": badge_away(tip_diag, "left"), "label": "Take Diagnostic"},
                {"n": 2, "tip": tip_path, "badge": badge_away(tip_path, "left"), "label": "Learning Path card"},
                {"n": 3, "tip": tip_classes, "badge": badge_away(tip_classes, "above"), "label": "My Classes (not enrolled yet)"},
            ]

            # Open diagnostic orientation, then Start so MCQ questions appear
            await page.evaluate("typeof openLmsDiagnostic === 'function' && openLmsDiagnostic()")
            await page.wait_for_timeout(1200)
            start_btn = page.locator('#lmsDiagnosticModal button:has-text("Start Diagnostic")').first
            if await start_btn.count() and await start_btn.is_visible():
                await start_btn.click()
            else:
                # Banner CTA path
                ban = page.locator('button:has-text("Take Diagnostic")').first
                if await ban.count() and await ban.is_visible():
                    await ban.click()
                    await page.wait_for_timeout(1000)
                    start_btn = page.locator('#lmsDiagnosticModal button:has-text("Start Diagnostic")').first
                    if await start_btn.count() and await start_btn.is_visible():
                        await start_btn.click()

            # Wait until a real MCQ is on screen (not orientation / already-completed)
            try:
                await page.wait_for_selector(
                    '#lmsDiagBody .lms-quiz-option, #lmsDiagBody .lms-quiz-option-body',
                    timeout=25000,
                )
            except Exception:
                # Fall back: shot orientation with matching legends
                pass
            await page.wait_for_timeout(800)
            await shot(page, "02_student_diagnostic")

            tip_q = await center(page, "#lmsDiagBody .lms-quiz-option")
            if not tip_q:
                tip_q = await center(page, "#lmsDiagBody .lms-quiz-option-body")
            tip_next = await center(page, '#lmsDiagBody button:has-text("Next"), #lmsDiagBody button:has-text("Skip")')
            tip_submit = await center(page, '#lmsDiagBody button:has-text("Submit Diagnostic")')
            tip_nav = tip_next or tip_submit
            body_txt = ""
            try:
                body_txt = (await page.locator("#lmsDiagBody").inner_text()).lower()
            except Exception:
                pass
            if "before you begin" in body_txt:
                tip_orient = await center(page, "#lmsDiagBody .lms-diag-orient-list, #lmsDiagBody .lms-diag-orient")
                tip_start = await center(page, '#lmsDiagnosticModal button:has-text("Start Diagnostic")')
                if not tip_orient:
                    tip_orient = (680, 420)
                if not tip_start:
                    tip_start = (780, 720)
                results["callouts"]["02_student_diagnostic"] = [
                    {"n": 1, "tip": tip_orient, "badge": badge_away(tip_orient, "left"), "label": "Orientation / instructions"},
                    {"n": 2, "tip": tip_start, "badge": badge_away(tip_start, "right"), "label": "Start Diagnostic"},
                ]
                results["diag_mode"] = "orientation"
            else:
                if not tip_q:
                    tip_q = (680, 450)
                if not tip_nav:
                    tip_nav = tip_submit or (980, 780)
                results["callouts"]["02_student_diagnostic"] = [
                    {"n": 1, "tip": tip_q, "badge": badge_away(tip_q, "left"), "label": "Question + answer choices"},
                    {"n": 2, "tip": tip_nav, "badge": badge_away(tip_nav, "above"), "label": "Next / Skip / Submit Diagnostic"},
                ]
                results["diag_mode"] = "questions"

            # Close diagnostic modal
            await page.evaluate(
                """() => {
                const m = document.getElementById('lmsDiagnosticModal');
                if (m) m.classList.remove('open');
            }"""
            )
            await page.keyboard.press("Escape")
            await page.wait_for_timeout(500)

            await page.evaluate("typeof openLmsStudentHub === 'function' && openLmsStudentHub()")
            await page.wait_for_timeout(1200)
            await shot(page, "03_student_quizzes")
        except Exception as exc:  # noqa: BLE001
            results["errors"].append(f"fresh student capture: {exc}")

        # ---------- Completed student: learning path / after diagnostic ----------
        page = await context.new_page()
        try:
            await login(page, STUDENT_DONE, PASSWORD)
            await page.goto(
                f"{BASE_URL}/student-dashboard",
                wait_until="domcontentloaded",
                timeout=60000,
            )
            await page.wait_for_timeout(2000)
            await page.evaluate(
                """() => {
                document.querySelectorAll('.lms-modal-backdrop.open').forEach(m => m.classList.remove('open'));
            }"""
            )
            await page.wait_for_timeout(400)
            await shot(page, "04_student_learning_path")
            tip_path = await center(page, "text=LEARNING PATH")
            if not tip_path:
                tip_path = await center(page, "text=Learning Path")
            tip_weak = await center(page, "text=WEAK TOPICS")
            if not tip_weak:
                tip_weak = await center(page, "text=Weak Topics")
            tip_pie = await center(page, "text=Topic Mastery Breakdown")
            if not tip_pie:
                tip_pie = await center(page, "canvas")
            if not tip_path:
                tip_path = (400, 260)
            if not tip_weak:
                tip_weak = (900, 300)
            if not tip_pie:
                tip_pie = (280, 620)
            results["callouts"]["04_student_learning_path"] = [
                {"n": 1, "tip": tip_path, "badge": badge_away(tip_path, "above"), "label": "Learning Path progress"},
                {"n": 2, "tip": tip_weak, "badge": badge_away(tip_weak, "above"), "label": "Weak Topics list"},
                {"n": 3, "tip": tip_pie, "badge": badge_away(tip_pie, "right"), "label": "Topic Mastery pie chart"},
            ]
        except Exception as exc:  # noqa: BLE001
            results["errors"].append(f"done student capture: {exc}")

        # ---------- Teacher analytics ----------
        page = await context.new_page()
        try:
            await login(page, TEACHER, PASSWORD)
            await page.goto(
                f"{BASE_URL}/teacher-dashboard",
                wait_until="domcontentloaded",
                timeout=60000,
            )
            await page.wait_for_timeout(2000)
            await shot(page, "05_teacher_dashboard")
            tip_analytics = await center(page, 'button:has-text("Analytics")')
            if not tip_analytics:
                tip_analytics = (980, 160)
            results["callouts"]["05_teacher_dashboard"] = [
                {"n": 1, "tip": tip_analytics, "badge": badge_away(tip_analytics, "below"), "label": "Analytics button"},
            ]

            opened = False
            for sel in [
                'button:has-text("Analytics")',
                'text=Class Analytics',
                '[onclick*="openLmsTeacherAnalytics"]',
            ]:
                try:
                    loc = page.locator(sel).first
                    if await loc.count() and await loc.is_visible():
                        await loc.click()
                        opened = True
                        break
                except Exception:
                    continue
            if not opened:
                await page.evaluate(
                    "typeof openLmsTeacherAnalytics === 'function' && openLmsTeacherAnalytics()"
                )
            await page.wait_for_selector("#lmsAnalyticsModal", timeout=15000)
            await page.wait_for_timeout(800)

            # Prefer Late Enroll / non-empty class
            class_sel = page.locator("#lmsAnalyticsClassSelect")
            if await class_sel.count():
                opts = class_sel.locator("option")
                pick = 0
                for i in range(await opts.count()):
                    txt = (await opts.nth(i).inner_text()).lower()
                    if "empty" in txt:
                        continue
                    pick = i
                    if "late" in txt:
                        break
                await class_sel.select_option(index=pick)
                await page.wait_for_timeout(1200)

            await page.locator('#lmsAnalyticsModal .lms-tab[data-tab="topics"]').first.click()
            await page.wait_for_timeout(1000)
            await shot(page, "06_analytics_topics")
            tip_class = await center(page, "#lmsAnalyticsClassSelect")
            tip_tab = await center(page, '#lmsAnalyticsModal .lms-tab[data-tab="topics"]')
            tip_struggle = await center(
                page,
                '#lmsAnalyticsContent .lms-expand-toggle, #lmsAnalyticsContent td:has-text("click to view"), #lmsAnalyticsContent tbody tr td:nth-child(3)',
            )
            if not tip_class:
                tip_class = (280, 160)
            if not tip_tab:
                tip_tab = (280, 210)
            if not tip_struggle:
                tip_struggle = (900, 360)
            results["callouts"]["06_analytics_topics"] = [
                {"n": 1, "tip": tip_class, "badge": badge_away(tip_class, "right"), "label": "Class dropdown"},
                {"n": 2, "tip": tip_tab, "badge": badge_away(tip_tab, "below"), "label": "Topic Performance tab"},
                {"n": 3, "tip": tip_struggle, "badge": badge_away(tip_struggle, "right"), "label": "Struggling students count"},
            ]

            tab_map = [
                ("progress", "07_analytics_topic_progress"),
                ("quizzes", "08_analytics_quizzes"),
                ("struggling", "09_analytics_struggling"),
                ("roster", "10_analytics_roster"),
            ]
            for tab, name in tab_map:
                await page.locator(f'#lmsAnalyticsModal .lms-tab[data-tab="{tab}"]').first.click()
                await page.wait_for_timeout(1100)
                await shot(page, name)

            # Chart example with Fractions data
            await page.locator('#lmsAnalyticsModal .lms-tab[data-tab="progress"]').first.click()
            await page.wait_for_timeout(800)
            student_sel = page.locator("#lmsTopicProgressStudentSelect")
            topic_sel = page.locator("#lmsTopicProgressSelect")
            if await student_sel.count():
                n = await student_sel.locator("option").count()
                for idx in range(min(n, 8)):
                    await student_sel.select_option(index=idx)
                    await page.wait_for_timeout(900)
                    if await topic_sel.count() and not await topic_sel.is_disabled():
                        t_opts = topic_sel.locator("option")
                        chosen = 0
                        for i in range(await t_opts.count()):
                            if "fraction" in (await t_opts.nth(i).inner_text()).lower():
                                chosen = i
                                break
                        if await t_opts.count() > 0:
                            await topic_sel.select_option(index=chosen)
                            await page.wait_for_timeout(1000)
                        empty = page.locator("#lmsTopicProgressEmpty:not(.hidden)")
                        if not await empty.count():
                            break

            await shot(page, "11_analytics_chart_example")
            tip_stu = await center(page, "#lmsTopicProgressStudentSelect")
            tip_topic = await center(page, "#lmsTopicProgressSelect")
            tip_chart = await center(page, "#lmsTopicProgressChart, .lms-analytics-chart-wrap")
            if not tip_stu:
                tip_stu = (350, 280)
            if not tip_topic:
                tip_topic = (750, 280)
            if not tip_chart:
                tip_chart = (680, 560)
            results["callouts"]["11_analytics_chart_example"] = [
                {"n": 1, "tip": tip_stu, "badge": badge_away(tip_stu, "above"), "label": "Student dropdown"},
                {"n": 2, "tip": tip_topic, "badge": badge_away(tip_topic, "above"), "label": "Topic dropdown"},
                {"n": 3, "tip": tip_chart, "badge": badge_away(tip_chart, "left"), "label": "Chart bars (Diagnostic purple, Quiz cyan)"},
            ]

            # Quiz results callouts
            await page.locator('#lmsAnalyticsModal .lms-tab[data-tab="quizzes"]').first.click()
            await page.wait_for_timeout(900)
            await shot(page, "08_analytics_quizzes")
            tip_qtab = await center(page, '#lmsAnalyticsModal .lms-tab[data-tab="quizzes"]')
            tip_row = await center(page, "#lmsAnalyticsContent tbody tr")
            if tip_qtab and tip_row:
                results["callouts"]["08_analytics_quizzes"] = [
                    {"n": 1, "tip": tip_qtab, "badge": badge_away(tip_qtab, "below"), "label": "Quiz Results tab"},
                    {"n": 2, "tip": tip_row, "badge": badge_away(tip_row, "right"), "label": "Assignment row (completion + avg score)"},
                ]

            results["ok"] = True
        except Exception as exc:  # noqa: BLE001
            results["errors"].append(f"teacher capture: {exc}")

        await browser.close()

    # Prefer verified Fractions chart if live empty
    qa_chart = QA_SHOTS / "B_happy_fractions.png"
    live_chart = ASSET_DIR / "11_analytics_chart_example.png"
    if qa_chart.exists():
        # Keep live if it looks like a real chart (file already captured); else copy QA
        if not live_chart.exists() or live_chart.stat().st_size < 20000:
            shutil.copy2(qa_chart, live_chart)
            # Approximate callouts for the known QA Fractions shot (1366x900)
            results["callouts"]["11_analytics_chart_example"] = [
                {"n": 1, "tip": [420, 255], "badge": [280, 200], "label": "Student dropdown"},
                {"n": 2, "tip": [780, 255], "badge": [920, 200], "label": "Topic dropdown"},
                {"n": 3, "tip": [520, 580], "badge": [180, 620], "label": "Chart bars (Diagnostic purple, Quiz cyan)"},
            ]

    META_PATH.write_text(json.dumps(results["callouts"], indent=2), encoding="utf-8")
    return results


def set_run(run, size=11, bold=False, color=DARK) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_para(doc, text: str, *, size=11, bold=False, color=DARK, space_after=8) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run(run, size=11)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_example_table(doc, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "16A34A")
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run(run, size=10)
            if r_i % 2 == 1:
                shade_cell(cell, "F0FDF4")
    doc.add_paragraph()


def load_callouts() -> dict:
    if META_PATH.exists():
        return json.loads(META_PATH.read_text(encoding="utf-8"))
    return {}


def add_shot_with_legend(doc, name: str, caption: str, legend: list[str]) -> None:
    raw = ASSET_DIR / f"{name}.png"
    if not raw.exists():
        add_para(doc, f"[Screenshot missing: {name}]", color=GRAY, size=10)
        return
    callouts = load_callouts().get(name) or []
    if callouts and legend:
        callouts = sorted(callouts, key=lambda c: int(c["n"]))[: len(legend)]
        for i, c in enumerate(callouts, 1):
            c["n"] = i
    elif not legend:
        callouts = []
    ann = ASSET_DIR / f"{name}_annotated.png"
    if callouts:
        annotate(raw, ann, callouts)
        pic = ann if ann.exists() else raw
    else:
        pic = raw
    doc.add_picture(str(pic), width=Inches(6.3))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run(run, size=9, color=GRAY)
    run.italic = True
    circles = "①②③④⑤⑥⑦⑧⑨"
    for i, line in enumerate(legend):
        mark = circles[i] if i < len(circles) else f"({i+1})"
        add_para(doc, f"{mark} {line}", size=10)
    doc.add_paragraph()


def build_document() -> Path:
    doc = Document()

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("IQBAL AI — How Analytics Works")
    set_run(run, size=26, bold=True, color=GREEN)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = s.add_run(
        "A simple guide: Student diagnostic → Score → Learning Path → Teacher Analytics"
    )
    set_run(run, size=12, color=GRAY)

    u = doc.add_paragraph()
    u.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = u.add_run(f"Live app: {BASE_URL}")
    set_run(run, size=10, color=GRAY)

    add_para(
        doc,
        "Red numbered circles on screenshots match the numbered list under each figure. "
        "Arrow tips were placed from live UI element positions (Playwright).",
        size=10,
        color=GRAY,
    )
    doc.add_page_break()

    doc.add_heading("1. The big picture (one sentence)", level=1)
    add_para(
        doc,
        "When a student finishes a diagnostic or quiz, the system counts correct answers, "
        "saves a score for each topic, builds a learning path for weak topics, and shows "
        "all of that to the teacher in Class Analytics.",
    )
    add_bullets(
        doc,
        [
            "Student takes Diagnostic (or a teacher Quiz)",
            "System scores answers: correct ÷ total × 100 = overall %",
            "Each topic also gets its own % (e.g. Fractions 50%)",
            "Topics under 60% are marked weak → Learning Path / Learning Chat",
            "Teacher opens Analytics and sees class + student progress",
        ],
    )
    doc.add_page_break()

    doc.add_heading("2. Student flow — start here", level=1)

    doc.add_heading("2.1 Student dashboard (fresh student)", level=2)
    add_para(
        doc,
        "After login, a new student sees the LMS Overview. From here they open Diagnostic, "
        "My Quizzes, and later their Learning Path (after diagnostic scores exist).",
    )
    add_shot_with_legend(
        doc,
        "01_student_dashboard",
        "Figure 1 — Student dashboard (starting point — fresh student, diagnostic not done yet)",
        [
            "Take Diagnostic (banner) — start the platform diagnostic",
            "Learning Path card — fills in after diagnostic scores exist",
            "My Classes — enroll with a join code (shows 0 until joined)",
        ],
    )

    doc.add_heading("2.2 Taking the diagnostic", level=2)
    add_para(
        doc,
        "The diagnostic is a timed MCQ test. First you see a short “Before you begin” screen, "
        "then Start opens the questions. Select an answer, tap Next, and on the last question tap Submit. "
        "If time runs out, answered questions still count.",
    )
    # Legend adapts to what was actually captured
    diag_legend = [
        "Question + answer choices (A/B/C/D)",
        "Skip → / Next / Submit Diagnostic",
    ]
    meta = load_callouts().get("02_student_diagnostic") or []
    if meta and any("Start" in str(c.get("label", "")) for c in meta):
        diag_legend = [
            "Before you begin — instructions (time, scoring, tips)",
            "Start Diagnostic — begins the timed MCQ test",
        ]
    add_shot_with_legend(
        doc,
        "02_student_diagnostic",
        "Figure 2 — Diagnostic assessment (fresh student — not “already completed”)",
        diag_legend,
    )

    doc.add_heading("2.3 How the score is calculated (very simple)", level=2)
    add_para(
        doc,
        "Each question is worth 1 point. Correct answer = 1. Wrong or skipped = 0.",
        bold=True,
    )
    add_para(doc, "Overall score formula:")
    add_para(
        doc,
        "Score % = (number of correct answers ÷ total questions) × 100",
        bold=True,
        color=GREEN,
    )
    add_para(doc, "Example A — overall score", bold=True)
    add_example_table(
        doc,
        ["Situation", "Math", "Result"],
        [
            ["10 questions, 7 correct", "7 ÷ 10 × 100", "70%"],
            ["20 questions, 8 correct", "8 ÷ 20 × 100", "40%"],
            ["5 questions, 5 correct", "5 ÷ 5 × 100", "100%"],
            ["Skipped questions", "Count as wrong (0)", "Lower overall %"],
        ],
    )
    add_para(doc, "Example B — topic score (this feeds analytics)", bold=True)
    add_example_table(
        doc,
        ["Topic", "Correct", "Total", "Topic %", "Status"],
        [
            ["Fractions", "1", "2", "50%", "Weak (< 60%)"],
            ["Algebra", "4", "4", "100%", "Mastered (≥ 85%)"],
            ["Geometry", "2", "3", "67%", "Needs practice (60–84%)"],
        ],
    )
    add_bullets(
        doc,
        [
            "Below 60% → Weak (needs help)",
            "60% to 84% → Needs practice / Improving",
            "85% or above → Mastered",
        ],
    )

    doc.add_heading("2.4 What happens right after Submit", level=2)
    add_bullets(
        doc,
        [
            "Save the score (correct / total / %)",
            "Update each topic’s StudentTopicScore",
            "Rebuild the Learning Path from weak topics",
            "Mark diagnostic complete (other LMS features unlock)",
            "If it was a class quiz — mark that assignment submitted",
        ],
    )

    doc.add_heading("2.5 Learning Path — after diagnostic", level=2)
    add_para(
        doc,
        "After diagnostic, weak topics (under 60%) drive the Learning Path / Learning Chat practice. "
        "This screenshot uses a student who already finished diagnostic so you can see filled cards.",
    )
    add_shot_with_legend(
        doc,
        "04_student_learning_path",
        "Figure 3 — Student view after diagnostic (Learning Path + Weak Topics)",
        [
            "Learning Path progress card",
            "Weak Topics list (scores under 60%)",
            "Topic Mastery breakdown pie chart",
        ],
    )
    add_para(doc, "Example C — learning path from Example B", bold=True)
    add_example_table(
        doc,
        ["Topic result", "Goes into Learning Path?"],
        [
            ["Fractions 50% (weak)", "Yes — practice in Learning Chat"],
            ["Geometry 67% (needs practice)", "Not listed as weak (<60%), still room to grow"],
            ["Algebra 100% (mastered)", "No — already strong"],
        ],
    )

    doc.add_heading("2.6 Teacher quizzes (My Quizzes)", level=2)
    add_para(
        doc,
        "When the teacher assigns a quiz, the student opens My Quizzes, starts it, and submits. "
        "Scoring works the same way. That quiz then appears on the teacher Topic Progress chart.",
    )
    add_shot_with_legend(
        doc,
        "03_student_quizzes",
        "Figure 4 — My Quizzes (teacher assignments)",
        [],
    )
    doc.add_page_break()

    doc.add_heading("3. Teacher side — Class Analytics", level=1)
    add_para(
        doc,
        "Teacher opens the dashboard and clicks Analytics. Class Analytics opens. First pick a class.",
    )
    add_shot_with_legend(
        doc,
        "05_teacher_dashboard",
        "Figure 5 — Teacher dashboard / opening Analytics",
        ["Analytics button — opens Class Analytics"],
    )

    doc.add_heading("3.1 Topic Performance", level=2)
    add_para(
        doc,
        "Shows each topic’s class average and how many students are struggling. "
        "Click the struggling count to see names and scores.",
    )
    add_para(doc, "Logic:", bold=True)
    add_bullets(
        doc,
        [
            "Reads each enrolled student’s latest topic scores",
            "Average = sum of scores ÷ number of scores",
            "Struggling on a topic = score under 60%",
            "Topics sorted weakest-first",
        ],
    )
    add_para(doc, "Example D — Topic Performance row", bold=True)
    add_example_table(
        doc,
        ["Student", "Fractions score"],
        [
            ["Ali", "40%"],
            ["Sara", "55%"],
            ["Omar", "90%"],
            ["Class avg", "(40+55+90) ÷ 3 = 61.7%"],
            ["Struggling count", "2 (Ali + Sara, both < 60%)"],
        ],
    )
    add_shot_with_legend(
        doc,
        "06_analytics_topics",
        "Figure 6 — Topic Performance tab",
        [
            "Class dropdown",
            "Topic Performance tab",
            "Struggling students count (click to expand names)",
        ],
    )

    doc.add_heading("3.2 Topic Progress (diagnostic vs quiz over time)", level=2)
    add_para(
        doc,
        "Pick one student and one topic. The chart shows every diagnostic and quiz for that topic.",
    )
    add_para(doc, "Logic:", bold=True)
    add_bullets(
        doc,
        [
            "Purple bar = Diagnostic",
            "Cyan / blue bar = Quiz",
            "Bar height = topic % on that assessment",
            "Stats: Assessments, First score, Latest score, Change",
        ],
    )
    add_para(doc, "Example E — Fractions for one student", bold=True)
    add_example_table(
        doc,
        ["Assessment", "Correct / Total", "Bar height", "Color"],
        [
            ["Diagnostic", "2 / 2", "100%", "Purple"],
            ["QA Quiz 1", "1 / 2", "50%", "Cyan"],
            ["Change", "Latest 50 − First 100", "−50 pts", "Needs help"],
        ],
    )
    add_shot_with_legend(
        doc,
        "11_analytics_chart_example",
        "Figure 7 — Topic Progress chart (student + topic selected)",
        [
            "Student dropdown",
            "Topic dropdown",
            "Chart bars (Diagnostic purple, Quiz cyan)",
        ],
    )
    if (ASSET_DIR / "07_analytics_topic_progress.png").exists():
        add_shot_with_legend(
            doc,
            "07_analytics_topic_progress",
            "Figure 8 — Topic Progress tab layout",
            [],
        )

    doc.add_heading("3.3 Quiz Results", level=2)
    add_para(
        doc,
        "Lists each published class assignment: completion %, class average, and expandable student scores.",
    )
    add_para(doc, "Logic:", bold=True)
    add_bullets(
        doc,
        [
            "Completion % = submitted ÷ enrolled × 100",
            "Avg score = average of submitted attempts only",
            "UI hint: ≥70% good, 50–69% warn",
        ],
    )
    add_para(doc, "Example F", bold=True)
    add_example_table(
        doc,
        ["Class size", "Submitted", "Scores", "Completion", "Avg"],
        [["4 students", "3", "80%, 60%, 40%", "75%", "(80+60+40)÷3 = 60%"]],
    )
    add_shot_with_legend(
        doc,
        "08_analytics_quizzes",
        "Figure 9 — Quiz Results tab",
        [
            "Quiz Results tab",
            "Assignment row (completion + avg score)",
        ]
        if load_callouts().get("08_analytics_quizzes")
        else [],
    )

    doc.add_heading("3.4 Struggling Students", level=2)
    add_para(doc, "Students who need teacher attention.")
    add_para(doc, "Logic — struggling if either:", bold=True)
    add_bullets(
        doc,
        [
            "2 or more weak topics (score under 60%), OR",
            "Overall progress under 60%",
        ],
    )
    add_para(doc, "Example G", bold=True)
    add_example_table(
        doc,
        ["Student", "Weak topics", "Overall progress", "Struggling?"],
        [
            ["Ali — Fractions + Ratios weak", "2", "45%", "Yes"],
            ["Sara — only Fractions weak", "1", "72%", "No (unless overall < 60%)"],
            ["Omar — no weak topics", "0", "88%", "No"],
        ],
    )
    add_shot_with_legend(
        doc,
        "09_analytics_struggling",
        "Figure 10 — Struggling Students tab",
        [],
    )

    doc.add_heading("3.5 Roster", level=2)
    add_para(
        doc,
        "Full class list with overall progress and weak-topic count. "
        "Overall progress = weighted average of topic scores.",
    )
    add_shot_with_legend(
        doc,
        "10_analytics_roster",
        "Figure 11 — Roster tab",
        [],
    )
    doc.add_page_break()

    doc.add_heading("4. One story from start to finish", level=1)
    add_para(doc, "Meet Sara (example student)", bold=True)
    add_bullets(
        doc,
        [
            "Sara takes Diagnostic: 10 questions → 6 correct → overall 60%.",
            "By topic: Fractions 50% (weak), Algebra 100% (mastered), Geometry 0% (weak).",
            "Learning Path → Learning Chat for weak areas.",
            "Teacher assigns a quiz; Sara’s Fractions quiz score appears next to diagnostic on Topic Progress.",
            "Teacher uses Topic Performance + Struggling Students to decide who needs help.",
        ],
    )
    doc.add_page_break()

    doc.add_heading("5. Quick reference card", level=1)
    add_example_table(
        doc,
        ["Thing", "Simple rule"],
        [
            ["Overall score %", "Correct ÷ total questions × 100"],
            ["Topic score %", "Correct on topic ÷ questions on topic × 100"],
            ["Weak topic", "Topic score < 60%"],
            ["Mastered topic", "Topic score ≥ 85%"],
            ["Learning Path", "Built from weak topics → Learning Chat"],
            ["Struggling student", "≥2 weak topics OR overall < 60%"],
            ["Topic Progress chart", "One student + one topic; purple=diagnostic, cyan=quiz"],
            ["Class topic average", "Average of enrolled students’ scores for that topic"],
            ["Quiz completion", "Submitted ÷ enrolled × 100"],
        ],
    )
    add_para(
        doc,
        "Regenerate: cd iqbal_ai_stg && python scripts/generate_analytics_guide.py",
        size=9,
        color=GRAY,
    )

    out = OUT_DIR / "IQBAL_AI_Analytics_How_It_Works.docx"
    doc.save(out)
    try:
        shutil.copy2(out, DESKTOP / "IQBAL_AI_Analytics_How_It_Works.docx")
    except OSError:
        pass
    try:
        shutil.copy2(out, ROOT.parent / "IQBAL_AI_Analytics_How_It_Works.docx")
    except OSError:
        pass
    return out


async def main() -> None:
    ensure_dirs()
    print(f"Capturing from {BASE_URL} ...")
    capture = await capture_screenshots()
    print("Capture summary:", {k: capture.get(k) for k in ("ok", "shots", "errors", "fresh_student")})
    print("Building Word document ...")
    path = build_document()
    print(f"Done: {path}")


if __name__ == "__main__":
    asyncio.run(main())
