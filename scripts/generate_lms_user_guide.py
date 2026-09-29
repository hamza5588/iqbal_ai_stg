#!/usr/bin/env python3
"""Generate IQBAL AI LMS end-to-end user guide (Word + PDF) with annotated UI screenshots."""

from __future__ import annotations

import os
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "guides" / "assets"
OUTPUT_DIR = ROOT / "docs" / "guides"
BASE_URL = os.environ.get("GUIDE_BASE_URL", "https://209.23.10.34")
ADMIN_USER = os.environ.get("GUIDE_ADMIN_USER", "admin@iqbalai.com")
ADMIN_PASS = os.environ.get("GUIDE_ADMIN_PASS", "admin123")
ADMIN_EMAIL_FALLBACKS = [
    ADMIN_USER,
    "admin@iqbal.edu",
    "admin@iqbalai.com",
]
TEACHER_USER = os.environ.get("GUIDE_TEACHER_USER", "teacher@iqbalai.com")
TEACHER_PASS = os.environ.get("GUIDE_TEACHER_PASS", "teacher123")
STUDENT_USER = os.environ.get("GUIDE_STUDENT_USER", "student@iqbalai.com")
STUDENT_PASS = os.environ.get("GUIDE_STUDENT_PASS", "student123")


@dataclass
class Annotation:
    x: int
    y: int
    label: str
    color: str = "#E11D48"


@dataclass
class GuideSection:
    title: str
    role: str
    overview: str
    steps: list[str]
    screenshot: str
    annotations: list[Annotation]


def ensure_dirs() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def annotate_image(src: Path, dst: Path, annotations: list[Annotation]) -> None:
    if not src.exists():
        return
    img = Image.open(src).convert("RGBA")
    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    try:
        font = ImageFont.truetype("arial.ttf", 18)
        font_sm = ImageFont.truetype("arial.ttf", 14)
    except OSError:
        font = ImageFont.load_default()
        font_sm = font

    for i, ann in enumerate(annotations):
        x, y = ann.x, ann.y
        color = ann.color
        # Arrow line from label box to point
        label_w = min(280, max(120, len(ann.label) * 8))
        box_x = max(10, min(x - label_w // 2, img.width - label_w - 10))
        box_y = max(10, y - 70 - (i % 3) * 8)
        draw.line([(box_x + label_w // 2, box_y + 34), (x, y)], fill=color, width=3)
        draw.ellipse([x - 7, y - 7, x + 7, y + 7], fill=color, outline="white", width=2)
        draw.rounded_rectangle(
            [box_x, box_y, box_x + label_w, box_y + 34],
            radius=8,
            fill=(255, 255, 255, 235),
            outline=color,
            width=2,
        )
        draw.text((box_x + 8, box_y + 8), ann.label, fill="#111827", font=font_sm)

    out = Image.alpha_composite(img, overlay).convert("RGB")
    out.save(dst, quality=92)


def placeholder(path: Path, title: str, subtitle: str) -> None:
    w, h = 1280, 720
    img = Image.new("RGB", (w, h), "#F8FAFC")
    draw = ImageDraw.Draw(img)
    try:
        font_lg = ImageFont.truetype("arial.ttf", 42)
        font_md = ImageFont.truetype("arial.ttf", 24)
    except OSError:
        font_lg = ImageFont.load_default()
        font_md = font_lg
    draw.rounded_rectangle([40, 40, w - 40, h - 40], radius=20, outline="#16A34A", width=4)
    draw.text((80, 120), title, fill="#14532D", font=font_lg)
    draw.text((80, 200), subtitle, fill="#475569", font=font_md)
    draw.text((80, h - 100), f"URL: {BASE_URL}", fill="#64748B", font=font_md)
    img.save(path)


async def capture_screenshots() -> None:
    try:
        from playwright.async_api import async_playwright
    except ImportError:
        print("Playwright not installed; using placeholder images.")
        return

    async def login(page, email: str, password: str) -> None:
        await page.goto(f"{BASE_URL}/auth/login", wait_until="domcontentloaded", timeout=60000)
        await page.fill('input[name="useremail"], input#email', email)
        await page.fill('input[name="password"], input#password', password)
        await page.click('button[type="submit"]')
        await page.wait_for_timeout(3000)

    async def shot(page, name: str) -> None:
        await page.wait_for_timeout(800)
        await page.screenshot(path=str(ASSETS / f"{name}.png"), full_page=False)

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        context = await browser.new_context(
            viewport={"width": 1440, "height": 900},
            ignore_https_errors=True,
        )

        # Admin
        page = await context.new_page()
        admin_ok = False
        for admin_email in ADMIN_EMAIL_FALLBACKS:
            try:
                await login(page, admin_email, ADMIN_PASS)
                await page.goto(f"{BASE_URL}/admin/", wait_until="domcontentloaded", timeout=60000)
                if "/admin" not in page.url:
                    continue
                admin_ok = True
                break
            except Exception:
                continue
        try:
            if not admin_ok:
                raise RuntimeError("admin login failed")
            await page.wait_for_timeout(1500)
            for sel in [
                'text=Diagnostic Assessment',
                'a:has-text("Diagnostic Assessment")',
                'span:has-text("Diagnostic Assessment")',
            ]:
                try:
                    await page.click(sel, timeout=3000)
                    break
                except Exception:
                    continue
            await page.evaluate("document.getElementById('diagnostic-section')?.scrollIntoView()")
            await page.wait_for_timeout(1000)
            await shot(page, "admin_diagnostic_section")
        except Exception as exc:
            print(f"Admin screenshot failed: {exc}")
            placeholder(ASSETS / "admin_diagnostic_section.png", "Admin — Diagnostic Assessment", "Upload Q&A PDF + Target PDFs")

        # Teacher
        page = await context.new_page()
        try:
            await login(page, TEACHER_USER, TEACHER_PASS)
            await page.goto(f"{BASE_URL}/teacher-dashboard", wait_until="networkidle", timeout=60000)
            await shot(page, "teacher_dashboard")
            for fab, name in [
                ("Analytics", "teacher_analytics"),
                ("PDF Quiz", "teacher_pdf_quiz"),
                ("Classes", "teacher_classes"),
                ("AI Tutor", "teacher_ai_tutor"),
            ]:
                try:
                    await page.click(f'button:has-text("{fab}")', timeout=5000)
                    await shot(page, name)
                    await page.keyboard.press("Escape")
                    await page.wait_for_timeout(500)
                except Exception:
                    pass
            try:
                await page.click('button:has-text("Assign Quiz")', timeout=5000)
                await shot(page, "teacher_assign_quiz")
            except Exception:
                pass
        except Exception as exc:
            print(f"Teacher screenshot failed: {exc}")
            for name, title in [
                ("teacher_dashboard", "Teacher Dashboard"),
                ("teacher_analytics", "Class Analytics"),
                ("teacher_pdf_quiz", "PDF Quiz Builder"),
                ("teacher_assign_quiz", "Assign Quiz"),
                ("teacher_classes", "Manage Classes"),
                ("teacher_ai_tutor", "AI Tutor"),
            ]:
                placeholder(ASSETS / f"{name}.png", title, "Teacher LMS feature")

        # Student
        page = await context.new_page()
        try:
            await login(page, STUDENT_USER, STUDENT_PASS)
            await page.goto(f"{BASE_URL}/student-dashboard", wait_until="networkidle", timeout=60000)
            await shot(page, "student_dashboard")
            for fab, name in [
                ("Diagnostic", "student_diagnostic"),
                ("My Quizzes", "student_quizzes"),
                ("Join Class", "student_join_class"),
            ]:
                try:
                    await page.click(f'button:has-text("{fab}")', timeout=5000)
                    await shot(page, name)
                    await page.keyboard.press("Escape")
                    await page.wait_for_timeout(500)
                except Exception:
                    pass
        except Exception as exc:
            print(f"Student screenshot failed: {exc}")
            for name, title in [
                ("student_dashboard", "Student Dashboard"),
                ("student_diagnostic", "Diagnostic Assessment"),
                ("student_quizzes", "My Quizzes"),
                ("student_join_class", "Join Class"),
            ]:
                placeholder(ASSETS / f"{name}.png", title, "Student LMS feature")

        await browser.close()


def build_sections() -> list[GuideSection]:
    return [
        GuideSection(
            title="Admin: Upload & Publish Platform Diagnostic",
            role="Admin",
            overview="Admin ek baar platform-wide diagnostic publish karta hai. Is mein Diagnostic Q&A PDF aur Target content PDFs shamil hain.",
            steps=[
                "Login karein: /auth/login → Admin dashboard (/admin/)",
                "Sidebar se 'Diagnostic Assessment' select karein",
                "Diagnostic Q&A PDF upload karein (questions + answer key)",
                "Ek ya zyada Target content PDFs add karein (Learning Chat ke liye)",
                "'Upload & Publish Diagnostic' par click karein",
                "Progress bar complete hone ke baad diagnostic auto-publish ho jata hai",
            ],
            screenshot="admin_diagnostic_section",
            annotations=[
                Annotation(180, 220, "Diagnostic Q&A PDF"),
                Annotation(520, 320, "Target content PDFs"),
                Annotation(900, 180, "Upload & Publish"),
                Annotation(700, 520, "Active diagnostic status"),
            ],
        ),
        GuideSection(
            title="Teacher: Class Analytics",
            role="Teacher",
            overview="Teacher har class ka performance dekh sakta hai — topics, quiz results, struggling students, aur roster.",
            steps=[
                "Teacher dashboard par 'Analytics' FAB button click karein",
                "Class dropdown se class select karein",
                "Topic Performance: weak students count par click karke names dekhein",
                "Quiz Results: har assignment ke student scores expand karein",
                "Struggling Students: weak topics count par click karke topic list dekhein",
                "Roster: overall progress aur status dekhein",
            ],
            screenshot="teacher_analytics",
            annotations=[
                Annotation(200, 120, "Select class"),
                Annotation(350, 170, "Topic Performance tab"),
                Annotation(500, 170, "Quiz Results tab"),
                Annotation(680, 170, "Struggling Students"),
                Annotation(400, 420, "Click count to expand"),
            ],
        ),
        GuideSection(
            title="Teacher: PDF Quiz Builder",
            role="Teacher",
            overview="Teacher apni PDF se MCQ quiz generate karta hai aur publish karta hai.",
            steps=[
                "'PDF Quiz' FAB button click karein",
                "Q&A PDF upload karein",
                "'Upload & Generate MCQs' — pipeline MCQs banata hai",
                "Preview questions check karein",
                "'Publish Quiz' — ab ye Assign Quiz mein available hoga",
            ],
            screenshot="teacher_pdf_quiz",
            annotations=[
                Annotation(400, 280, "Upload PDF"),
                Annotation(650, 360, "Generate MCQs"),
                Annotation(500, 520, "Preview questions"),
                Annotation(750, 600, "Publish Quiz"),
            ],
        ),
        GuideSection(
            title="Teacher: Assign Quiz to Class",
            role="Teacher",
            overview="Published quiz ko class ke students ko assignment ke tor par bhejna.",
            steps=[
                "'Assign Quiz' FAB click karein",
                "Assignment title, class, aur published quiz select karein",
                "Optional due date set karein",
                "'Create & Publish' — students ko 'My Quizzes' mein dikhega",
            ],
            screenshot="teacher_assign_quiz",
            annotations=[
                Annotation(350, 250, "Assignment title"),
                Annotation(350, 330, "Select class"),
                Annotation(350, 410, "Select quiz"),
                Annotation(350, 520, "Create & Publish"),
            ],
        ),
        GuideSection(
            title="Teacher: Manage Classes",
            role="Teacher",
            overview="Classes create karna, join code share karna, aur students manage karna.",
            steps=[
                "'Classes' FAB click karein",
                "Nayi class create karein (name + grade)",
                "Join code students ko share karein",
                "'Manage' se roster dekhein ya students manually add karein",
            ],
            screenshot="teacher_classes",
            annotations=[
                Annotation(300, 200, "Create class form"),
                Annotation(500, 380, "Join code"),
                Annotation(750, 420, "Student names list"),
                Annotation(900, 300, "Manage button"),
            ],
        ),
        GuideSection(
            title="Teacher: AI Tutor",
            role="Teacher",
            overview="Teaching assistant chat — pedagogy, question design, aur classroom help.",
            steps=[
                "'AI Tutor' FAB click karein",
                "Apna sawal type karein",
                "AI step-by-step guidance deta hai (direct answer nahi)",
            ],
            screenshot="teacher_ai_tutor",
            annotations=[
                Annotation(700, 200, "Chat messages"),
                Annotation(500, 750, "Type question"),
                Annotation(900, 750, "Send"),
            ],
        ),
        GuideSection(
            title="Student: Diagnostic Assessment",
            role="Student",
            overview="Platform diagnostic ek martaba liya jata hai. Timer ke sath MCQs, phir weak/strong topics result.",
            steps=[
                "'Diagnostic' FAB ya onboarding se start karein",
                "Har question ka jawab select karein → Next",
                "Back button left side par hai — peechle sawal par ja sakte hain",
                "Last question par 'Submit Diagnostic'",
                "Results: score, weak topics, aur 'Start Learning Chat' option",
            ],
            screenshot="student_diagnostic",
            annotations=[
                Annotation(200, 650, "Back (left)"),
                Annotation(900, 650, "Next / Submit"),
                Annotation(500, 300, "Question & options"),
                Annotation(500, 120, "Timer / progress"),
            ],
        ),
        GuideSection(
            title="Student: My Quizzes (Teacher Assignments)",
            role="Student",
            overview="Teacher ki assigned quizzes yahan dikhti hain.",
            steps=[
                "'My Quizzes' FAB click karein",
                "Assignment par 'Start' click karein",
                "Har question answer karke Next karein",
                "Submit par score aur learning path update hota hai",
            ],
            screenshot="student_quizzes",
            annotations=[
                Annotation(400, 300, "Assignment list"),
                Annotation(750, 350, "Start button"),
                Annotation(500, 500, "Quiz taking area"),
            ],
        ),
        GuideSection(
            title="Student: Learning Chat (After Diagnostic)",
            role="Student",
            overview="Diagnostic ke baad weak topics par PDF-grounded Learning Chat — tutor hints deta hai, direct answer nahi.",
            steps=[
                "Diagnostic submit ke baad 'Start Learning Chat' click karein",
                "Ya Learning Path se weak topic step open karein",
                "Har MCQ ka jawab submit karein",
                "'Ask Tutor' / 'Need more help' se guided hints lein",
                "'Pause & Exit' se baad mein resume kar sakte hain",
            ],
            screenshot="student_dashboard",
            annotations=[
                Annotation(150, 750, "Learning Path"),
                Annotation(400, 750, "Diagnostic FAB"),
                Annotation(650, 750, "My Quizzes"),
                Annotation(500, 200, "LMS Overview"),
            ],
        ),
        GuideSection(
            title="Student: Join Class",
            role="Student",
            overview="Teacher ka diya hua join code se class mein enroll hona.",
            steps=[
                "'Join Class' FAB click karein",
                "Join code enter karein (teacher se milega)",
                "Join — ab teacher ki assignments aur analytics mein show honge",
            ],
            screenshot="student_join_class",
            annotations=[
                Annotation(500, 350, "Join code input"),
                Annotation(500, 450, "Join button"),
            ],
        ),
    ]


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("IQBAL AI — LMS User Guide\n")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x16, 0xA3, 0x4A)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "End-to-End Guide: Admin Diagnostic • Teacher (Analytics, PDF Quiz, AI Tutor) • Student Diagnostic\n"
    ).font.size = Pt(14)
    doc.add_paragraph(f"Application URL: {BASE_URL}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_toc(doc: Document, sections: list[GuideSection]) -> None:
    doc.add_heading("Table of Contents", level=1)
    for i, sec in enumerate(sections, 1):
        doc.add_paragraph(f"{i}. [{sec.role}] {sec.title}", style="List Number")
    doc.add_page_break()


def add_flow_overview(doc: Document) -> None:
    doc.add_heading("How the LMS Flow Works", level=1)
    doc.add_paragraph(
        "1. Admin platform diagnostic publish karta hai (Q&A PDF + target content PDFs).\n"
        "2. Student diagnostic leta hai — weak topics identify hote hain.\n"
        "3. Student Learning Chat se weak areas par practice karta hai.\n"
        "4. Teacher PDF se quiz banata hai aur class ko assign karta hai.\n"
        "5. Teacher Analytics se class performance monitor karta hai."
    )
    doc.add_page_break()


def build_word(sections: list[GuideSection]) -> Path:
    doc = Document()
    add_cover(doc)
    add_toc(doc, sections)
    add_flow_overview(doc)

    for sec in sections:
        doc.add_heading(f"[{sec.role}] {sec.title}", level=1)
        doc.add_heading("Overview", level=2)
        doc.add_paragraph(sec.overview)
        doc.add_heading("Step-by-Step", level=2)
        for step in sec.steps:
            doc.add_paragraph(step, style="List Number")

        raw = ASSETS / f"{sec.screenshot}.png"
        ann = ASSETS / f"{sec.screenshot}_annotated.png"
        if raw.exists():
            annotate_image(raw, ann, sec.annotations)
            doc.add_heading("UI Screenshot (annotated)", level=2)
            if ann.exists():
                doc.add_picture(str(ann), width=Inches(6.5))
            else:
                doc.add_picture(str(raw), width=Inches(6.5))
            cap = doc.add_paragraph("Red arrows = important UI controls")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cap.runs:
                cap.runs[0].italic = True
        else:
            doc.add_paragraph("[Screenshot not captured — run with Playwright on live server]")
        doc.add_page_break()

    out_docx = OUTPUT_DIR / "IQBAL_AI_LMS_User_Guide.docx"
    doc.save(out_docx)
    return out_docx


def build_pdf_reportlab(sections: list[GuideSection], docx_path: Path) -> Optional[Path]:
    """Fallback PDF when docx2pdf / MS Word unavailable."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import inch
        from reportlab.pdfgen import canvas
    except ImportError:
        return None

    pdf_path = OUTPUT_DIR / "IQBAL_AI_LMS_User_Guide.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=A4)
    w, h = A4
    y = h - 0.75 * inch

    def new_page_if_needed(need: float = 60) -> None:
        nonlocal y
        if y < need:
            c.showPage()
            y = h - 0.75 * inch

    c.setFont("Helvetica-Bold", 20)
    c.drawString(0.75 * inch, y, "IQBAL AI — LMS User Guide")
    y -= 28
    c.setFont("Helvetica", 11)
    c.drawString(0.75 * inch, y, f"URL: {BASE_URL}")
    c.showPage()
    y = h - 0.75 * inch

    for sec in sections:
        c.setFont("Helvetica-Bold", 14)
        new_page_if_needed(100)
        c.drawString(0.75 * inch, y, f"[{sec.role}] {sec.title[:70]}")
        y -= 22
        c.setFont("Helvetica-Bold", 11)
        c.drawString(0.75 * inch, y, "Overview")
        y -= 16
        c.setFont("Helvetica", 10)
        for line in _wrap(sec.overview, 95):
            new_page_if_needed()
            c.drawString(0.85 * inch, y, line)
            y -= 14
        y -= 6
        c.setFont("Helvetica-Bold", 11)
        c.drawString(0.75 * inch, y, "Steps")
        y -= 16
        c.setFont("Helvetica", 10)
        for i, step in enumerate(sec.steps, 1):
            for line in _wrap(f"{i}. {step}", 95):
                new_page_if_needed()
                c.drawString(0.85 * inch, y, line)
                y -= 14
        ann = ASSETS / f"{sec.screenshot}_annotated.png"
        raw = ASSETS / f"{sec.screenshot}.png"
        img_path = ann if ann.exists() else raw
        if img_path.exists():
            new_page_if_needed(200)
            c.drawImage(str(img_path), 0.75 * inch, y - 4.5 * inch, width=6.5 * inch, height=4.0 * inch, preserveAspectRatio=True, anchor='nw')
            y -= 4.6 * inch
        c.showPage()
        y = h - 0.75 * inch

    c.save()
    return pdf_path


def _wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur: list[str] = []
    for word in words:
        trial = (" ".join(cur + [word])).strip()
        if len(trial) <= width:
            cur.append(word)
        else:
            if cur:
                lines.append(" ".join(cur))
            cur = [word]
    if cur:
        lines.append(" ".join(cur))
    return lines or [""]


def build_pdf(docx_path: Path, sections: list[GuideSection]) -> Optional[Path]:
    pdf_path = OUTPUT_DIR / "IQBAL_AI_LMS_User_Guide.pdf"
    try:
        from docx2pdf import convert

        convert(str(docx_path), str(pdf_path))
        return pdf_path
    except Exception as exc:
        print(f"docx2pdf unavailable ({exc}); using ReportLab fallback.")
        return build_pdf_reportlab(sections, docx_path)


def main() -> int:
    ensure_dirs()
    import asyncio

    asyncio.run(capture_screenshots())

    # Ensure placeholders for any missing shots
    for sec in build_sections():
        p = ASSETS / f"{sec.screenshot}.png"
        if not p.exists():
            placeholder(p, sec.title, sec.role)

    sections = build_sections()
    docx_path = build_word(sections)
    pdf_path = build_pdf(docx_path, sections)
    print(f"Word guide: {docx_path}")
    if pdf_path:
        print(f"PDF guide:  {pdf_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
