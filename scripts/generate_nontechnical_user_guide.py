#!/usr/bin/env python3
"""Build a non-technical IQBALAI Word user guide from the provided screenshots."""

from __future__ import annotations

import math
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "guides"
ASSET_DIR = OUT_DIR / "assets" / "nontechnical_guide"
CURSOR_ASSETS = Path(
    r"C:\Users\user\.cursor\projects\c-Users-user-Desktop-iqbalai-v1-1\assets"
)
DESKTOP_COPY = Path(r"C:\Users\user\Desktop\IQBALAI_User_Guide.docx")
PROJECT_COPY = ROOT.parent / "IQBALAI_User_Guide.docx"

GREEN = RGBColor(0x16, 0xA3, 0x4A)
DARK = RGBColor(0x11, 0x18, 0x27)
GRAY = RGBColor(0x47, 0x55, 0x69)
RED = RGBColor(0xDC, 0x26, 0x26)

SRC = {
    "admin_active": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-b5a942bd-4a47-499d-aca8-0c7777e3d630.png",
    "admin_cards": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-645a170a-5f76-4d51-8cbe-3c07f12a1ce7.png",
    "admin_upload": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-fb810e05-7afc-4419-8cb5-d1b555f8f32a.png",
    "teacher_classes": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-6137c1bc-36c1-4e92-a5c6-d7129c18d3a2.png",
    "teacher_pdf_quiz": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-e894299e-d599-46dd-8ce7-c32865310e63.png",
    "teacher_analytics": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-af4282e2-1728-4b9d-bcb5-cf2e774d080a.png",
    "teacher_tutor": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-0883cbef-0889-4604-88a0-05171683faee.png",
    "student_home": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-45daab2a-a8c0-4432-a414-3d5978972979.png",
    "student_quizzes": CURSOR_ASSETS
    / "c__Users_user_AppData_Roaming_Cursor_User_workspaceStorage_dcd9e19689df97890af86435f58a806c_images_image-962e428d-06a9-4358-8610-4e41a837aa92.png",
}


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    for folder in (Path(r"C:\Windows\Fonts"), Path("/usr/share/fonts")):
        path = folder / name
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def arrow_head(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str, size: int = 16) -> None:
    ang = math.atan2(end[1] - start[1], end[0] - start[0])
    p1 = (
        end[0] - size * math.cos(ang - 0.45),
        end[1] - size * math.sin(ang - 0.45),
    )
    p2 = (
        end[0] - size * math.cos(ang + 0.45),
        end[1] - size * math.sin(ang + 0.45),
    )
    draw.polygon([end, p1, p2], fill=color)


def annotate(src: Path, dst: Path, callouts: list[dict]) -> Path:
    img = Image.open(src).convert("RGBA")
    overlay = Image.new("RGBA", img.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    badge_font = font(18, bold=True)
    color = "#DC2626"

    for item in callouts:
        n = item["n"]
        tx, ty = item["tip"]
        bx, by = item["badge"]
        r = 16
        draw.line([(bx, by), (tx, ty)], fill=color, width=4)
        arrow_head(draw, (bx, by), (tx, ty), color, size=15)
        draw.ellipse([bx - r, by - r, bx + r, by + r], fill=color, outline="white", width=3)
        text = str(n)
        bbox = draw.textbbox((0, 0), text, font=badge_font)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text((bx - tw / 2, by - th / 2 - 1), text, fill="white", font=badge_font)
        draw.ellipse([tx - 6, ty - 6, tx + 6, ty + 6], fill=color, outline="white", width=2)

    out = Image.alpha_composite(img, overlay).convert("RGB")
    dst.parent.mkdir(parents=True, exist_ok=True)
    out.save(dst, quality=94)
    return dst


def set_run(run, size=11, bold=False, color=DARK, name="Calibri") -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_border(cell) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "E5E7EB")
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def add_para(doc, text, size=11, bold=False, color=DARK, space_after=8, center=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def add_feature_table(doc, rows: list[tuple[str, str, str, str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    headers = ["#", "What you see", "What it is for", "How to use it"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run(run, size=10, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(cell, "166534")
        set_cell_border(cell)
    for r_idx, row in enumerate(rows, start=1):
        bg = "F0FDF4" if r_idx % 2 else "FFFFFF"
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_run(run, size=10, bold=(c_idx == 0), color=DARK)
            shade_cell(cell, bg)
            set_cell_border(cell)
    doc.add_paragraph()


def add_steps(doc, steps: list[str]) -> None:
    add_para(doc, "Do this, in order:", size=12, bold=True, color=GREEN, space_after=4)
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.15)
        run = p.add_run(f"Step {i}.  ")
        set_run(run, size=11, bold=True, color=GREEN)
        run2 = p.add_run(step)
        set_run(run2, size=11, color=DARK)
    doc.add_paragraph()


def add_screenshot(doc, path: Path, width: float = 6.5) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    p.add_run().add_picture(str(path), width=Inches(width))


def add_note(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Tip: ")
    set_run(run, size=11, bold=True, color=RGBColor(0xB4, 0x53, 0x09))
    run2 = p.add_run(text)
    set_run(run2, size=11, color=GRAY)


def build_images() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    specs = {
        "admin_active": [
            {"n": 1, "tip": (80, 370), "badge": (40, 250)},
            {"n": 2, "tip": (948, 28), "badge": (800, 85)},
            {"n": 3, "tip": (500, 175), "badge": (780, 90)},
            {"n": 4, "tip": (930, 265), "badge": (980, 175)},
            {"n": 5, "tip": (300, 310), "badge": (180, 210)},
            {"n": 6, "tip": (900, 342), "badge": (980, 410)},
        ],
        "admin_cards": [
            {"n": 1, "tip": (90, 28), "badge": (40, 70)},
            {"n": 2, "tip": (220, 52), "badge": (520, 20)},
            {"n": 3, "tip": (120, 95), "badge": (40, 150)},
            {"n": 4, "tip": (960, 115), "badge": (980, 50)},
            {"n": 5, "tip": (280, 500), "badge": (40, 430)},
        ],
        "admin_upload": [
            {"n": 1, "tip": (180, 40), "badge": (40, 90)},
            {"n": 2, "tip": (280, 75), "badge": (560, 30)},
            {"n": 3, "tip": (200, 470), "badge": (40, 420)},
            {"n": 4, "tip": (90, 530), "badge": (280, 500)},
            {"n": 5, "tip": (115, 590), "badge": (320, 560)},
        ],
        "teacher_classes": [
            {"n": 1, "tip": (70, 28), "badge": (40, 85)},
            {"n": 2, "tip": (880, 28), "badge": (980, 75)},
            {"n": 3, "tip": (990, 400), "badge": (980, 300)},
            {"n": 4, "tip": (420, 95), "badge": (200, 40)},
            {"n": 5, "tip": (350, 200), "badge": (80, 200)},
            {"n": 6, "tip": (450, 320), "badge": (80, 340)},
            {"n": 7, "tip": (640, 375), "badge": (820, 330)},
        ],
        "teacher_pdf_quiz": [
            {"n": 1, "tip": (140, 28), "badge": (40, 80)},
            {"n": 2, "tip": (790, 22), "badge": (760, 80)},
            {"n": 3, "tip": (400, 145), "badge": (700, 120)},
            {"n": 4, "tip": (90, 230), "badge": (280, 200)},
            {"n": 5, "tip": (170, 325), "badge": (400, 330)},
        ],
        "teacher_analytics": [
            {"n": 1, "tip": (950, 22), "badge": (900, 70)},
            {"n": 2, "tip": (180, 95), "badge": (40, 70)},
            {"n": 3, "tip": (160, 150), "badge": (40, 200)},
            {"n": 4, "tip": (200, 260), "badge": (40, 320)},
            {"n": 5, "tip": (720, 300), "badge": (900, 220)},
            {"n": 6, "tip": (400, 550), "badge": (720, 575)},
        ],
        "teacher_tutor": [
            {"n": 1, "tip": (200, 105), "badge": (80, 40)},
            {"n": 2, "tip": (740, 105), "badge": (880, 50)},
            {"n": 3, "tip": (450, 250), "badge": (160, 200)},
            {"n": 4, "tip": (400, 430), "badge": (160, 520)},
            {"n": 5, "tip": (706, 438), "badge": (860, 520)},
        ],
        "student_home": [
            {"n": 1, "tip": (500, 42), "badge": (280, 80)},
            {"n": 2, "tip": (975, 72), "badge": (850, 20)},
            {"n": 3, "tip": (150, 120), "badge": (40, 70)},
            {"n": 4, "tip": (580, 145), "badge": (750, 70)},
            {"n": 5, "tip": (70, 234), "badge": (40, 180)},
            {"n": 6, "tip": (180, 300), "badge": (40, 280)},
            {"n": 7, "tip": (500, 360), "badge": (720, 330)},
            {"n": 8, "tip": (990, 430), "badge": (880, 370)},
        ],
        "student_quizzes": [
            {"n": 1, "tip": (160, 115), "badge": (40, 40)},
            {"n": 2, "tip": (920, 115), "badge": (980, 40)},
            {"n": 3, "tip": (120, 200), "badge": (40, 170)},
            {"n": 4, "tip": (150, 230), "badge": (40, 290)},
            {"n": 5, "tip": (840, 220), "badge": (960, 170)},
        ],
    }
    paths = {}
    for key, callouts in specs.items():
        src = SRC[key]
        if not src.exists():
            raise FileNotFoundError(src)
        raw = ASSET_DIR / f"{key}_raw.png"
        shutil.copy2(src, raw)
        paths[key] = annotate(src, ASSET_DIR / f"{key}_annotated.png", callouts)
    return paths


def build_doc(images: dict[str, Path]) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    add_para(doc, "IQBALAI", size=14, bold=True, color=GREEN, center=True, space_after=2)
    add_para(doc, "How to Use the Application", size=28, bold=True, color=DARK, center=True, space_after=6)
    add_para(
        doc,
        "A simple picture guide for non-technical people",
        size=14,
        color=GRAY,
        center=True,
        space_after=6,
    )
    add_para(
        doc,
        "Red numbered arrows on every screenshot show you exactly where to look.\nMatch the number in the picture with the table under it.",
        size=12,
        color=DARK,
        center=True,
        space_after=18,
    )

    add_para(doc, "Who this guide is for", size=16, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "IQBALAI is a learning app with three kinds of users. You only need to read the part that matches your role.",
        size=11,
        space_after=8,
    )
    add_feature_table(
        doc,
        [
            ("A", "Admin", "Sets up the one-time diagnostic test and study PDFs for the whole platform.", "Read Part A."),
            ("B", "Teacher", "Creates classes, makes quizzes from PDFs, and checks student progress.", "Read Part B."),
            ("C", "Student", "Joins a class, takes tests, and practices weak topics with AI.", "Read Part C."),
        ],
    )

    add_para(doc, "How the whole app works (the big picture)", size=16, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "1. Admin uploads one diagnostic test and the study PDFs.\n"
        "2. Student takes that diagnostic once. The app finds weak topics.\n"
        "3. Student practices those weak topics in Learning Chat.\n"
        "4. Teacher creates a class, shares a join code, and assigns quizzes.\n"
        "5. Teacher opens Analytics to see who needs extra help.",
        size=11,
        space_after=10,
    )
    add_note(
        doc,
        "Button names in this guide are written exactly as they appear on the screen. If you cannot find something, look for those same words.",
    )

    # ---------- PART A ----------
    doc.add_page_break()
    add_para(doc, "Part A — Admin", size=22, bold=True, color=GREEN, space_after=4)
    add_para(doc, "Diagnostic Assessment: set up the test students take once", size=14, bold=True, space_after=8)
    add_para(
        doc,
        "Open Admin Dashboard after you log in. Use the left menu to open Diagnostic Assessment. "
        "Here you publish one platform diagnostic. Students can take it only one time. "
        "The app sets the timer from each question’s difficulty.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["admin_active"])
    add_feature_table(
        doc,
        [
            ("1", "Left menu — Diagnostic Assessment", "Takes you to this page from other admin tools.", "Click Diagnostic Assessment in the sidebar. The selected item is highlighted in light blue."),
            ("2", "Logout", "Signs you out of the admin account.", "Click the red Logout button when you are finished."),
            ("3", "Yellow message box", "Tells you the current rule: a diagnostic is already active.", "Read it before you change anything. You can still add more study PDFs. To replace the whole test, remove the diagnostic first."),
            ("4", "Remove diagnostic", "Deletes the active test so you can upload a new Q&A test.", "Click only if you really want a new test. Students will no longer use this one."),
            ("5", "Target PDFs list", "These PDFs are the study material Learning Chat uses after the test.", "Check the file names. Students learn from every PDF listed here."),
            ("6", "Remove (next to a PDF)", "Removes one study file from this diagnostic.", "Click Remove beside a file you no longer want. The other files stay."),
        ],
    )
    add_steps(
        doc,
        [
            "Log in and open Admin Dashboard.",
            "Click Diagnostic Assessment in the left menu.",
            "If no diagnostic is active, upload the Q&A PDF (questions and answers) and publish it.",
            "Add study-material PDFs under Target PDFs so Learning Chat has content.",
            "Use Remove only when a file or the whole diagnostic should be replaced.",
        ],
    )
    add_note(doc, "published means students can take this test. archived means it is old and no longer the live test.")

    add_para(doc, "Reading each diagnostic card", size=14, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "As you scroll, every diagnostic appears as a white card. The top card is usually the live one. Older tests are marked archived.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["admin_cards"])
    add_feature_table(
        doc,
        [
            ("1", "Card title", "The name of that diagnostic test.", "Use the title to tell tests apart (for example a class name or a date)."),
            ("2", "Status line", "Shows archived or published, how many questions, and the timer.", "A published test with 8 questions and ~5 min timer means students get about 5 minutes."),
            ("3", "TARGET PDFS", "How many study files are attached to this test.", "The number in brackets is the file count."),
            ("4", "Remove next to a file", "Deletes that one PDF from the card.", "Click if the wrong file was uploaded."),
            ("5", "Orange warning", "This test has no study PDFs yet.", "Upload PDFs using the form at the bottom of the page. Learning Chat needs these files."),
        ],
    )

    add_para(doc, "Upload extra study PDFs", size=14, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "Scroll to the bottom of the Diagnostic Assessment page. This form adds more study material. "
        "Students use all uploaded target PDFs in Learning Chat — not only one file.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["admin_upload"])
    add_feature_table(
        doc,
        [
            ("1", "Older diagnostic cards", "Past tests that are archived.", "You can still attach PDFs to them, but students take the active published test."),
            ("2", "Orange ‘No target content PDFs…’", "That card is missing study files.", "Upload PDFs below if Learning Chat should use content for that diagnostic."),
            ("3", "Add Target Content PDFs heading", "This is the upload area.", "Stay in this section to add study files."),
            ("4", "Choose files", "Opens your computer’s file picker.", "Click Choose files, select one or more PDFs, then confirm."),
            ("5", "Upload Target PDFs", "Sends the selected files to the platform.", "Click this blue button after you choose files. Wait until the list updates."),
        ],
    )
    add_steps(
        doc,
        [
            "Scroll to Add Target Content PDFs.",
            "Click Choose files and pick PDF study notes from your computer.",
            "Click the blue Upload Target PDFs button.",
            "Confirm the new files appear in the Target PDFs list on the active diagnostic.",
        ],
    )
    add_note(doc, "You can select more than one PDF at a time: hold Ctrl on Windows (or Cmd on Mac) while clicking files.")

    # ---------- PART B ----------
    doc.add_page_break()
    add_para(doc, "Part B — Teacher", size=22, bold=True, color=GREEN, space_after=4)
    add_para(doc, "Manage Classes: create a class and share the join code", size=14, bold=True, space_after=8)
    add_para(
        doc,
        "After you log in as a teacher you see your lessons. On the right, click Classes. "
        "A window opens where you set the grades you teach, create classes, and copy join codes for students.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["teacher_classes"])
    add_feature_table(
        doc,
        [
            ("1", "IQBALAI logo and top menu", "My Lessons, Set Prompt, and Chat History.", "Use My Lessons to return to your lesson list."),
            ("2", "+ Create New Lesson", "Starts a new lesson.", "Click when you want to build a new lesson, not a class."),
            ("3", "Right-side shortcut buttons", "Shortcuts: Classes, Assign Quiz, PDF Quiz, Analytics, AI Tutor.", "These round buttons sit on the right. Classes opened this window. Use Assign Quiz, PDF Quiz, Analytics, and AI Tutor for the other teacher tools."),
            ("4", "Your Teaching Grades", "Tells the app which grades you teach.", "Type grade numbers (for example 8,10), then click Save Teaching Grades."),
            ("5", "Create New Class", "Makes a new student group.", "Type a class name (example: Math 8A), pick Grade level, then click Create Class."),
            ("6", "Join code", "The password students type to enter your class.", "Share this code in class or by message. Each class has its own code."),
            ("7", "Manage", "Opens that class’s student list and settings.", "Click Manage to add or review students."),
        ],
    )
    add_steps(
        doc,
        [
            "Click Classes on the right.",
            "Save the grades you teach if they are not already listed.",
            "Enter a class name, choose the grade, and click Create Class.",
            "Copy the Join code and give it to students.",
            "Ask students to click Join Class on their dashboard and type that code.",
        ],
    )
    add_note(doc, "If a class says No students enrolled yet, nobody has used the join code. Share the code again.")

    add_para(doc, "PDF → MCQ Quiz Builder: turn a PDF into a quiz", size=14, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "Click PDF Quiz on the right-side menu. This window turns a question-and-answer PDF into multiple-choice questions you can assign later.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["teacher_pdf_quiz"], width=5.8)
    add_feature_table(
        doc,
        [
            ("1", "Window title", "Confirms you are in the quiz builder.", "If you opened the wrong window, click the X and try PDF Quiz again."),
            ("2", "X (close)", "Closes the window without creating a quiz.", "Click X to go back to your dashboard."),
            ("3", "Quiz title", "The name students will see.", "Type a clear name, for example Chapter 5 Post Test."),
            ("4", "Choose file", "Selects the Q&A PDF from your computer.", "Click Choose file and pick the PDF that contains questions and answers. The words No file chosen will change to the file name."),
            ("5", "Upload & Generate MCQs", "Uploads the PDF and asks the app to build the quiz.", "Click the green button and wait. Then preview the questions and publish when they look correct."),
        ],
    )
    add_steps(
        doc,
        [
            "Click PDF Quiz.",
            "Type a quiz title.",
            "Click Choose file and select your Q&A PDF.",
            "Click Upload & Generate MCQs and wait for the questions to appear.",
            "Check the questions, then publish. After that, use Assign Quiz to send it to a class.",
        ],
    )

    add_para(doc, "Class Analytics: see who is struggling", size=14, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "Click Analytics. This report shows topic scores for one class. Red numbers mean students need help. Click those numbers to see names.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["teacher_analytics"])
    add_feature_table(
        doc,
        [
            ("1", "X (close)", "Closes Analytics and returns to the dashboard.", "Click X when you are done."),
            ("2", "Select class", "Chooses which class’s data to show.", "Open the dropdown and pick the class, for example chem10a (10th)."),
            ("3", "Tabs", "Four views: Topic Performance, Quiz Results, Struggling Students, Roster.", "Click a tab to switch. Start with Topic Performance."),
            ("4", "Topic and Avg Score", "Class average for that topic.", "A low percentage (for example 0% or 33%) means the class is weak on that topic."),
            ("5", "Red badge — click to view", "How many students are struggling on that row.", "Click 1 student · click to view to see names and scores."),
            ("6", "Yellow Insight box", "A short summary of the table.", "Read it first. It tells you how many topics have struggling students."),
        ],
    )
    add_steps(
        doc,
        [
            "Click Analytics.",
            "Select the class.",
            "Stay on Topic Performance, or open Quiz Results / Struggling Students / Roster.",
            "Look for low Avg Score values.",
            "Click the red student count to see who needs help.",
        ],
    )

    add_para(doc, "Teaching Assistant: ask the AI tutor", size=14, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "Click AI Tutor. This chat helps you with teaching ideas, explanations, and classroom questions for any subject.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["teacher_tutor"], width=6.0)
    add_feature_table(
        doc,
        [
            ("1", "Teaching Assistant title", "You are in the AI tutor window.", "Use this for lesson help, not for grading the whole class."),
            ("2", "X (close)", "Closes the chat.", "Click X to return to the dashboard. You can open it again anytime."),
            ("3", "Message area", "Shows the conversation.", "The first line is a greeting. Answers appear here after you send a question."),
            ("4", "Type your question…", "Where you write.", "Click inside the box and type, for example: Give me 5 warm-up questions on fractions."),
            ("5", "Send", "Sends your question to the tutor.", "Click Send (or press Enter if it works in your browser) and wait for the reply."),
        ],
    )
    add_steps(
        doc,
        [
            "Click AI Tutor on the right.",
            "Click the box that says Type your question…",
            "Type your question in plain language.",
            "Click Send and read the answer in the white area above.",
        ],
    )

    # ---------- PART C ----------
    doc.add_page_break()
    add_para(doc, "Part C — Student", size=22, bold=True, color=GREEN, space_after=4)
    add_para(doc, "Your home screen (My Lessons)", size=14, bold=True, space_after=8)
    add_para(
        doc,
        "After you log in as a student, this page is your home. It shows your progress, weak topics, classes, quizzes, and shortcuts to practice.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["student_home"])
    add_feature_table(
        doc,
        [
            ("1", "Top menu", "Chat with Lesson, My Lessons, and Today.", "Stay on My Lessons for this overview. Use Chat with Lesson when you want to talk about a lesson."),
            ("2", "Join Class and My Quizzes", "Join a teacher’s class, or open your quiz list.", "Join Class: enter the code your teacher gave you. My Quizzes: see assignments."),
            ("3", "Overall Progress", "How much of your work is done, as a percent.", "37% means you have completed a little more than one third so far."),
            ("4", "Weak Topics", "Subjects the diagnostic found you need to practice.", "Read the names (Fractions, Algebra, Geometry) and the scores. These are the topics to practice first."),
            ("5", "Open Learning Chat", "Starts practice on your weak areas using the study PDFs.", "Click the green Open Learning Chat button and follow the questions."),
            ("6", "Quiz History", "Scores from quizzes you already took.", "Use this to see if your scores are improving."),
            ("7", "My Lessons search", "Finds lessons by name. All topics filters the list.", "Type a word in the search box, or change All topics. If you see No lessons available yet, none are assigned yet."),
            ("8", "Diagnostic / Practice / Ask Tutor", "Quick buttons that stay in the corner.", "Diagnostic: take the one-time placement test. Practice: extra exercises. Ask Tutor: get help from the AI."),
        ],
    )
    add_steps(
        doc,
        [
            "If your teacher gave you a code, click Join Class and type it.",
            "Click Diagnostic (corner button) and take the test once. Use Next to move forward.",
            "After the test, read Weak Topics on this page.",
            "Click Open Learning Chat to practice those topics.",
            "Click My Quizzes when your teacher assigns a quiz.",
            "Use Ask Tutor if you are stuck.",
        ],
    )
    add_note(doc, "The diagnostic can be taken only once. Do not close the browser in the middle if you can avoid it. Pending Quizzes = 0 means you have no waiting tests.")

    add_para(doc, "Assignments & Quizzes: check a quiz you already submitted", size=14, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "Click My Quizzes. A window lists every assignment. You can see whether you submitted it and what score or progress you have.",
        size=11,
        space_after=8,
    )
    add_screenshot(doc, images["student_quizzes"])
    add_feature_table(
        doc,
        [
            ("1", "Assignments & Quizzes title", "This window is your quiz list.", "Open it from the green My Quizzes button."),
            ("2", "X (close)", "Closes the list and returns to the dashboard.", "Click X when you are finished looking."),
            ("3", "Quiz name", "The title your teacher gave the assignment.", "In this example the name is 1st q. Click the row if it offers Start or Open."),
            ("4", "Status: submitted", "You already turned this quiz in.", "submitted means the app received your answers. If it says not started, click to begin."),
            ("5", "Completed · 50%", "Your result or progress on that quiz.", "Here the score is 50%. Use Learning Chat to practice weak topics, then try the next quiz."),
        ],
    )
    add_steps(
        doc,
        [
            "Click My Quizzes.",
            "Find the quiz name your teacher told you about.",
            "If Status is not submitted, open it and answer the questions, then submit.",
            "If Status is submitted, read your percent on the right.",
            "Click X to go back to the dashboard.",
        ],
    )

    # ---------- Quick reference ----------
    doc.add_page_break()
    add_para(doc, "Quick reference", size=22, bold=True, color=GREEN, space_after=8)
    add_para(doc, "Colours and messages you will see", size=14, bold=True, space_after=6)
    add_feature_table(
        doc,
        [
            ("Green button", "Start, save, send, or create", "A safe next step.", "Click it to continue."),
            ("Blue button", "Upload files or choose a file", "Sends a PDF to the app.", "Choose the file first, then click upload."),
            ("Red Logout / Remove", "Sign out, or delete a test or PDF", "This cannot always be undone.", "Click only when you are sure."),
            ("Yellow box", "Helpful status message", "Explains what is allowed right now.", "Read it before you change settings."),
            ("Orange text", "Something is missing", "Usually no PDF is attached.", "Upload the missing PDF."),
            ("Red number badge", "Students who are struggling", "Needs the teacher’s attention.", "Click the number to see names."),
        ],
    )

    add_para(doc, "If something does not work", size=14, bold=True, color=GREEN, space_after=6)
    add_para(
        doc,
        "• Make sure you are logged in with the correct role (admin, teacher, or student).\n"
        "• Refresh the page and try again.\n"
        "• For file uploads, use a PDF only — not Word or photos.\n"
        "• Students: you need the teacher’s join code to appear in a class.\n"
        "• Teachers: publish a quiz before you assign it.\n"
        "• Admins: if a diagnostic is already active, remove it only when you want to replace the whole test.",
        size=11,
        space_after=12,
    )
    add_para(doc, "You are done with this guide. Open the app, find the matching screen, and follow the numbered arrows.", size=12, bold=True, color=DARK, center=True)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUT_DIR / "IQBALAI_User_Guide_Non_Technical.docx"
    doc.save(out_path)
    shutil.copy2(out_path, DESKTOP_COPY)
    shutil.copy2(out_path, PROJECT_COPY)
    return out_path


def main() -> None:
    images = build_images()
    path = build_doc(images)
    print(f"Wrote {path}")
    print(f"Copied to {DESKTOP_COPY}")
    print(f"Copied to {PROJECT_COPY}")


if __name__ == "__main__":
    main()
