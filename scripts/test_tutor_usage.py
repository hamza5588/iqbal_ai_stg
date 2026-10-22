"""Disposable-database tests for time accounting, report coverage and real XLSX output."""
import ast
import importlib.util
from datetime import datetime, timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('usage_test_models', ROOT/'app/models/database_models.py')
models = importlib.util.module_from_spec(spec)
spec.loader.exec_module(models)
tree = ast.parse((ROOT/'app/services/tutor_usage_service.py').read_text(encoding='utf-8'))
tree.body = [n for n in tree.body if not (isinstance(n, ast.ImportFrom) and (n.module or '').startswith('app.'))]
service = {name: getattr(models, name) for name in ['User', 'TutorUsageSession', 'TutorDailyUsage', 'SystemSettings']}
exec(compile(tree, 'tutor_usage_service.py', 'exec'), service)
NOW = datetime(2026, 10, 22, 10)


@pytest.fixture
def db():
    engine = create_engine('sqlite:///:memory:')
    models.Base.metadata.create_all(engine)
    with Session(engine) as db:
        for id_, role in [(1, 'student'), (2, 'student'), (3, 'teacher')]:
            db.add(models.User(id=id_, username=f'user{id_}', full_name=f'Student {id_}',
                               useremail=f'user{id_}@example.test', role=role, password='unused',
                               class_standard='9', medium='English', groq_api_key='',
                               school_name='School A' if id_ == 1 else 'School B', created_at=NOW-timedelta(days=90)))
        db.add(models.SystemSettings(key=service['START_KEY'], value=NOW.isoformat()))
        db.commit()
        yield db
    engine.dispose()


def heartbeat(db, token, seq, total, offset, user=1, now=NOW):
    return service['record_heartbeat'](db, user, {'session_id': token, 'seq': seq, 'total_ms': total}, now+timedelta(seconds=offset))


def test_active_time_retry_and_out_of_order(db):
    token = str(uuid4())
    assert heartbeat(db, token, 1, 0, 0) == 0
    assert heartbeat(db, token, 2, 15000, 15) == 15000
    assert heartbeat(db, token, 2, 15000, 20) == 0
    assert heartbeat(db, token, 1, 0, 20) == 0
    assert heartbeat(db, token, 3, 20000, 20) == 5000
    assert db.query(models.TutorDailyUsage).one().active_ms == 20000


def test_two_tabs_never_double_count(db):
    first, second = str(uuid4()), str(uuid4())
    for token in (first, second):
        heartbeat(db, token, 1, 0, 0)
    assert heartbeat(db, first, 2, 15000, 15) == 15000
    assert heartbeat(db, second, 2, 15000, 15) == 0
    assert heartbeat(db, second, 3, 30000, 30) == 15000
    assert heartbeat(db, first, 3, 30000, 30) == 0
    assert db.query(models.TutorDailyUsage).one().active_ms == 30000


def test_users_are_separate_and_session_cannot_be_stolen(db):
    first, second = str(uuid4()), str(uuid4())
    heartbeat(db, first, 1, 0, 0)
    with pytest.raises(PermissionError):
        heartbeat(db, first, 2, 15000, 15, user=2)
    heartbeat(db, second, 1, 0, 0, user=2)
    assert heartbeat(db, first, 2, 15000, 15) == 15000
    assert heartbeat(db, second, 2, 15000, 15, user=2) == 15000


def test_offline_gap_is_not_backfilled(db):
    token = str(uuid4())
    heartbeat(db, token, 1, 0, 0)
    assert heartbeat(db, token, 2, 100000, 100) == 0
    assert heartbeat(db, token, 3, 115000, 115) == 15000


def test_client_cannot_claim_more_than_server_elapsed_or_interval_cap(db):
    token = str(uuid4())
    heartbeat(db, token, 1, 0, 0)
    assert heartbeat(db, token, 2, 900000, 10) == 10000
    assert heartbeat(db, token, 3, 1800000, 50) == 30000


def test_idle_heartbeat_does_not_add_time(db):
    token = str(uuid4())
    heartbeat(db, token, 1, 0, 0)
    assert heartbeat(db, token, 2, 0, 15) == 0
    assert db.query(models.TutorDailyUsage).count() == 0


def test_split_at_pakistan_midnight(db):
    token = str(uuid4())
    before_midnight = datetime(2026, 10, 22, 18, 59, 50)
    heartbeat(db, token, 1, 0, 0, now=before_midnight)
    assert heartbeat(db, token, 2, 20000, 20, now=before_midnight) == 20000
    rows = db.query(models.TutorDailyUsage).order_by(models.TutorDailyUsage.usage_date).all()
    assert [(r.usage_date.isoformat(), r.active_ms) for r in rows] == [('2026-10-22', 10000), ('2026-10-23', 10000)]


@pytest.mark.parametrize('patch', [dict(seq=True), dict(total_ms=-1), dict(total_ms=2.5), dict(session_id='bad'), dict(user_id=2), dict(message='private text')])
def test_invalid_and_extra_payload_fields_rejected(db, patch):
    payload = {'session_id':str(uuid4()), 'seq':1, 'total_ms':0, **patch}
    with pytest.raises(ValueError):
        service['record_heartbeat'](db, 1, payload, NOW)
    assert db.query(models.TutorUsageSession).count() == 0


def test_teacher_and_deactivated_students_cannot_track(db):
    with pytest.raises(PermissionError):
        heartbeat(db, str(uuid4()), 1, 0, 0, user=3)
    db.get(models.User, 1).is_active = False
    with pytest.raises(PermissionError):
        heartbeat(db, str(uuid4()), 1, 0, 0)


def test_report_has_zero_users_coverage_filters_and_no_chat_fields(db):
    token = str(uuid4())
    heartbeat(db, token, 1, 0, 0)
    heartbeat(db, token, 2, 15000, 15)
    report = service['daily_report'](db, {'start':'2026-10-21','end':'2026-10-22'}, NOW)
    assert report['total'] == 4 and report['active_students'] == 1
    assert [r['active_seconds'] for r in report['rows']] == [None, None, 15, 0]
    assert [r['coverage'] for r in report['rows']] == ['Not tracked','Not tracked','Partial day','Partial day']
    assert set(report['rows'][0]) == {'date','student_id','student_name','school_name','active_seconds','coverage'}
    filtered = service['daily_report'](db, {'school':'School A'}, NOW)
    assert len(filtered['rows']) == 1 and filtered['total_active_seconds'] == 15
    assert service['daily_report'](db, {'search':'%'}, NOW)['total'] == 0


@pytest.mark.parametrize('params', [dict(start='bad'), dict(start='2026-10-23'), dict(start='2026-09-01',end='2026-10-22'), dict(start='2026-10-22',end='2026-10-21')])
def test_date_validation(db, params):
    with pytest.raises(ValueError):
        service['daily_report'](db, params, NOW)


def test_real_excel_numeric_time_and_formula_safe_names(db):
    db.get(models.User, 1).full_name = '=HYPERLINK("https://example.test")'
    token = str(uuid4())
    heartbeat(db, token, 1, 0, 0)
    heartbeat(db, token, 2, 15000, 15)
    report = service['daily_report'](db, {'school':'School A'}, NOW)
    workbook = load_workbook(service['build_workbook'](report))
    sheet = workbook['Daily Usage']
    assert sheet['C2'].data_type == 's' and sheet['C2'].value.startswith('=')
    assert sheet['E2'].value == timedelta(seconds=15)
    assert sheet['F2'].value == 0.25
    assert sheet.max_column == 7 and sheet.max_row == 2
    assert workbook.sheetnames == ['Daily Usage', 'Report Info']


def test_tracking_start_survives_repeated_initialization(db):
    db.commit()
    service['initialize_tracking'](db.get_bind())
    service['initialize_tracking'](db.get_bind())
    assert db.query(models.SystemSettings).filter_by(key=service['START_KEY']).one().value == NOW.isoformat()
