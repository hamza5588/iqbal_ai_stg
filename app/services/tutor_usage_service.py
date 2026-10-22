"""Active tutor time and daily reports. No dependency on chat history."""
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from uuid import UUID
from zoneinfo import ZoneInfo

from sqlalchemy import text

from app.models.database_models import User, TutorUsageSession, TutorDailyUsage, SystemSettings

REPORT_ZONE = ZoneInfo('Asia/Karachi')
START_KEY = 'tutor_usage_started_at'
MAX_INTERVAL_MS = 30_000
MAX_GAP_SECONDS = 45


def utc_now():
    return datetime.now(timezone.utc).replace(tzinfo=None)


def initialize_tracking(engine):
    """Run once at startup/migration; preserve the original coverage start on restart."""
    with engine.begin() as conn:
        conn.execute(text(
            "INSERT INTO system_settings (key, value, description) "
            "VALUES (:key, :value, 'Start of active AI Tutor usage tracking (UTC)') "
            "ON CONFLICT (key) DO NOTHING"
        ), {'key': START_KEY, 'value': utc_now().isoformat()})


def record_heartbeat(db, user_id, payload, now=None):
    now = now or utc_now()
    if not isinstance(payload, dict) or set(payload) != {'session_id', 'seq', 'total_ms'}:
        raise ValueError('Invalid usage update.')
    try:
        session_id = str(UUID(payload['session_id']))
    except (TypeError, ValueError, AttributeError):
        raise ValueError('Invalid usage session.') from None
    seq, total = payload['seq'], payload['total_ms']
    if type(seq) is not int or type(total) is not int or not (1 <= seq <= 2_147_483_647) or not (0 <= total <= 2_147_483_647):
        raise ValueError('Invalid usage counters.')
    # Serialize all devices for a user, including the first insert of a daily row.
    user = db.query(User).filter(User.id == user_id).with_for_update().first()
    if not user or user.role != 'student' or not user.is_active:
        raise PermissionError('Student access required.')
    cursor = db.get(TutorUsageSession, session_id)
    if cursor is None:
        db.add(TutorUsageSession(id=session_id, user_id=user_id, sequence=seq,
                                client_total_ms=total, last_seen_at=now))
        db.flush()
        return 0  # First contact establishes a server-time baseline.
    if cursor.user_id != user_id:
        raise PermissionError('Invalid usage session.')
    if seq <= cursor.sequence:
        return 0  # Retries and delayed requests cannot add time twice.
    if total < cursor.client_total_ms:
        raise ValueError('Usage counter cannot decrease.')
    gap = (now - cursor.last_seen_at).total_seconds()
    credit = min(total - cursor.client_total_ms, max(0, int(gap * 1000)), MAX_INTERVAL_MS)
    if gap < 0 or gap > MAX_GAP_SECONDS:
        credit = 0  # Offline/suspended intervals are never backfilled.
    cursor.sequence, cursor.client_total_ms, cursor.last_seen_at = seq, total, max(now, cursor.last_seen_at)
    start, credited = now - timedelta(milliseconds=credit), 0
    while start < now:
        local_day = start.replace(tzinfo=timezone.utc).astimezone(REPORT_ZONE).date()
        boundary = datetime.combine(local_day + timedelta(days=1), time.min, REPORT_ZONE).astimezone(timezone.utc).replace(tzinfo=None)
        end = min(now, boundary)
        daily = db.get(TutorDailyUsage, (user_id, local_day))
        if daily is None:
            daily = TutorDailyUsage(user_id=user_id, usage_date=local_day, active_ms=0, credited_until=start)
            db.add(daily)
        # Union overlapping device/tab intervals instead of summing them.
        effective_start = max(start, daily.credited_until)
        amount = max(0, int((end - effective_start).total_seconds() * 1000))
        daily.active_ms += amount
        daily.credited_until = max(daily.credited_until, end)
        credited += amount
        start = end
    db.flush()
    return credited


def parse_window(params, now=None):
    today = (now or utc_now()).replace(tzinfo=timezone.utc).astimezone(REPORT_ZONE).date()
    try:
        start = date.fromisoformat(params.get('start') or today.isoformat())
        end = date.fromisoformat(params.get('end') or start.isoformat())
    except (TypeError, ValueError):
        raise ValueError('Choose valid start and end dates.') from None
    if start > end or end > today or (end - start).days >= 31:
        raise ValueError('Choose up to 31 days, ending today or earlier.')
    return start, end


def daily_report(db, params, now=None):
    start, end = parse_window(params, now)
    setting = db.query(SystemSettings).filter(SystemSettings.key == START_KEY).first()
    started = datetime.fromisoformat(setting.value) if setting else None
    start_day = started.replace(tzinfo=timezone.utc).astimezone(REPORT_ZONE).date() if started else None
    query = db.query(User).filter(User.role == 'student')
    school = (params.get('school') or '').strip()
    if school:
        query = query.filter(User.school_name == school)
    search = (params.get('search') or '').strip()
    if len(search) > 255:
        raise ValueError('Search must be 255 characters or fewer.')
    if search:
        term = '%' + search.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_') + '%'
        query = query.filter(User.full_name.ilike(term, escape='\\') | User.username.ilike(term, escape='\\'))
    students = query.order_by(User.full_name, User.username, User.id).all()
    ids = [s.id for s in students]
    usage = db.query(TutorDailyUsage).filter(TutorDailyUsage.usage_date.between(start, end))
    # Filtering via a subquery avoids a parameter for every student in a large school.
    usage = usage.filter(TutorDailyUsage.user_id.in_(query.with_entities(User.id)))
    totals = {(u.user_id, u.usage_date): u.active_ms for u in usage}
    rows, active_ids, total_ms = [], set(), 0
    for offset in range((end - start).days + 1):
        day = start + timedelta(days=offset)
        for student in students:
            created_day = student.created_at.replace(tzinfo=timezone.utc).astimezone(REPORT_ZONE).date() if student.created_at else day
            if created_day > day:
                continue
            tracked = start_day is not None and day >= start_day
            ms = totals.get((student.id, day), 0) if tracked else None
            coverage = 'Not tracked' if not tracked else ('Partial day' if day == start_day else 'Tracked')
            rows.append({'date': day.isoformat(), 'student_id': student.id,
                         'student_name': student.full_name or student.username,
                         'school_name': student.school_name or 'Not provided',
                         'active_seconds': ms // 1000 if ms is not None else None,
                         'coverage': coverage})
            if ms:
                total_ms += ms
                active_ids.add(student.id)
    schools = [name for (name,) in db.query(User.school_name).filter(User.role == 'student', User.school_name.isnot(None), User.school_name != '').distinct().order_by(User.school_name)]
    return {'rows': rows, 'start': start.isoformat(), 'end': end.isoformat(), 'timezone': 'Asia/Karachi',
            'tracking_started_at': started.isoformat() + 'Z' if started else None,
            'schools': schools, 'total': len(rows), 'student_count': len(ids),
            'active_students': len(active_ids), 'total_active_seconds': total_ms // 1000}


def build_workbook(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Daily Usage'
    sheet.append(['Date', 'Student ID', 'Student Name', 'School Name', 'AI Tutor Time', 'Active Minutes', 'Coverage'])
    for row in report['rows']:
        seconds = row['active_seconds']
        sheet.append([date.fromisoformat(row['date']), row['student_id'], row['student_name'], row['school_name'],
                      timedelta(seconds=seconds) if seconds is not None else None,
                      round(seconds / 60, 2) if seconds is not None else None, row['coverage']])
        sheet.cell(sheet.max_row, 1).number_format = 'yyyy-mm-dd'
        sheet.cell(sheet.max_row, 5).number_format = '[h]:mm:ss'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in enumerate([14, 14, 30, 26, 20, 18, 18], start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1D4ED8')
    info = workbook.create_sheet('Report Info')
    for key, value in [('Timezone', report['timezone']), ('Start date', report['start']), ('End date', report['end']),
                       ('Tracking started (UTC)', report['tracking_started_at'] or 'Not started'),
                       ('Measurement', 'Visible, focused tutor panel after interaction; pauses after 2 minutes idle.'),
                       ('Coverage', 'The first tracking day is partial. Earlier dates are not tracked.'),
                       ('Offline', 'Offline time is excluded. Closing a disconnected tab may lose its last unsent interval.')]:
        info.append([key, value])
    info.column_dimensions['A'].width = 28
    info.column_dimensions['B'].width = 100
    # User-supplied names must remain text, including strings beginning with '='.
    for ws in workbook:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
